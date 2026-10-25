"""
Sample Test Dataset Generator.
Creates a realistic 'College_Documents' folder containing authentic PDF, DOCX, XLSX, and TXT files.
Matches user example specifications:
- Exam_Timetable.pdf
- Fee_Receipt.pdf
- Scholarship_Application.pdf
- Placement_Notice.pdf
- College_Circular.pdf
- Scholarship_Guidelines.docx
- Required_Documents_Checklist.txt
- Campus_Drive_Schedule.xlsx
- CS101_Syllabus.docx
- Tuition_Challan_Fall2024.pdf
"""

import os
from fpdf import FPDF
import docx
import openpyxl


def create_simple_pdf(filepath: str, title: str, content_lines: list):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, title, new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.ln(5)
    pdf.set_font("Helvetica", size=11)
    for line in content_lines:
        pdf.multi_cell(0, 7, line)
        pdf.ln(2)
    pdf.output(filepath)


def create_simple_docx(filepath: str, title: str, paragraphs: list):
    doc = docx.Document()
    doc.add_heading(title, level=1)
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.save(filepath)


def create_simple_xlsx(filepath: str, sheet_title: str, headers: list, rows: list):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(headers)
    for r in rows:
        ws.append(r)
    wb.save(filepath)


def create_simple_txt(filepath: str, content: str):
    with open(filepath, "w", encoding="utf-8") as f:
        f.write(content)


def generate_sample_college_documents(target_dir: str) -> str:
    """
    Creates realistic sample documents inside target_dir/College_Documents
    """
    sample_root = os.path.join(target_dir, "College_Documents")
    os.makedirs(sample_root, exist_ok=True)

    # 1. Exam_Timetable.pdf
    create_simple_pdf(
        os.path.join(sample_root, "Exam_Timetable.pdf"),
        "UNIVERSITY SEMESTER FINAL EXAMINATION TIMETABLE",
        [
            "Academic Year: 2024 - 2025 | Semester: 4 | Department of Computer Science",
            "Notice to all students: The end-term examinations are scheduled as follows:",
            "Course CS401 (Algorithms): Dec 10, 2024 (10:00 AM - 01:00 PM) - Hall A",
            "Course CS402 (Operating Systems): Dec 12, 2024 (10:00 AM - 01:00 PM) - Hall B",
            "Course CS403 (Database Systems): Dec 15, 2024 (10:00 AM - 01:00 PM) - Hall A",
            "Please bring your Hall Ticket / Admit Card and College ID Card. No digital devices allowed in examination hall."
        ]
    )

    # 2. Fee_Receipt.pdf
    create_simple_pdf(
        os.path.join(sample_root, "Fee_Receipt.pdf"),
        "OFFICIAL TUITION & COLLEGE FEE PAYMENT RECEIPT",
        [
            "Receipt No: REC-2024-88419 | Date of Payment: 14-Aug-2024",
            "Student Name: Alex Morgan | Enrollment ID: CS-2023-042",
            "Tuition Fee: $4,500.00",
            "Library & Laboratory Dues: $350.00",
            "Hostel & Dining Deposit: $1,200.00",
            "Total Paid: $6,050.00 (Status: COMPLETED / NO DUES)",
            "Transaction ID: TXN-993821094 | Payment Method: Online NetBanking"
        ]
    )

    # 3. Scholarship_Application.pdf
    create_simple_pdf(
        os.path.join(sample_root, "Scholarship_Application.pdf"),
        "MERIT-CUM-MEANS SCHOLARSHIP APPLICATION FORM",
        [
            "National Student Fellowship Scheme 2024-2025",
            "Applicant Name: Alex Morgan | Grade Point Average (GPA): 3.92 / 4.00",
            "Annual Family Income: Verified below threshold ($24,000)",
            "Applied Category: Merit & Low-Income Higher Education Grant",
            "Enclosed Documents: Income Certificate, Previous Semester Marksheet, Recommendation Letter.",
            "Declaration: I certify that all details submitted for this scholarship grant are accurate."
        ]
    )

    # 4. Placement_Notice.pdf
    create_simple_pdf(
        os.path.join(sample_root, "Placement_Notice.pdf"),
        "CAMPUS RECRUITMENT & PLACEMENT DRIVE NOTICE",
        [
            "Training and Placement Cell (T&P Cell) Announcement",
            "Company: Nexus Technologies Global Inc.",
            "Role: Associate Software Engineer (Full-Time CTC: $85,000 p.a.)",
            "Eligible Branches: B.Tech / B.E. Computer Science, Information Technology, Data Science",
            "Selection Process: Online Aptitude Test, Technical Interview, HR Interview",
            "Date of On-Campus Interview: November 22, 2024. Submit updated resume by Nov 15."
        ]
    )

    # 5. College_Circular.pdf
    create_simple_pdf(
        os.path.join(sample_root, "College_Circular.pdf"),
        "OFFICE OF THE REGISTRAR - GENERAL CIRCULAR",
        [
            "Circular Ref No: REG/CIR/2024/112 | Date: 05-Sep-2024",
            "Subject: Observance of National Science and Innovation Week",
            "All faculty, staff, and students are hereby informed that the annual science seminar",
            "and project exhibition will be organized from October 1st to October 5th.",
            "Attendance is mandatory for all undergraduate departments.",
            "By order of the Dean of Academic Affairs."
        ]
    )

    # 6. Scholarship_Guidelines.docx
    create_simple_docx(
        os.path.join(sample_root, "Scholarship_Guidelines.docx"),
        "Guidelines for Merit-cum-Means Financial Aid & Fellowship",
        [
            "1. Eligibility: Candidates must maintain a minimum CGPA of 3.5 without any active backlogs.",
            "2. Income Verification: Parents' or guardian's annual income tax certificate or Tahsildar income certificate is mandatory.",
            "3. Disbursement: Scholarship amounts will be credited directly to the student's registered bank account in two equal installments.",
            "4. Renewal Process: Students must re-apply at the beginning of each academic semester."
        ]
    )

    # 7. Required_Documents_Checklist.txt
    create_simple_txt(
        os.path.join(sample_root, "Required_Documents_Checklist.txt"),
        "CHECKLIST FOR SCHOLARSHIP AND FINANCIAL GRANT APPLICATION:\n"
        "1. Signed Scholarship Application Form\n"
        "2. Copy of Semester Fee Payment Receipt / Challan\n"
        "3. Official Grade Transcript / Marksheet\n"
        "4. Income Certificate from Competent Authority\n"
        "5. Government Identity Card (Aadhaar / Passport copy)\n"
        "6. Bank Account Passbook front page copy\n"
    )

    # 8. Campus_Drive_Schedule.xlsx
    create_simple_xlsx(
        os.path.join(sample_root, "Campus_Drive_Schedule.xlsx"),
        "Placements_2024",
        ["Company Name", "Industry", "Job Profile", "Interview Date", "Shortlisted Count", "CTC Offer"],
        [
            ["Nexus Technologies", "Software", "Associate Dev", "2024-11-22", 45, "$85,000"],
            ["FinTech Alpha", "Banking", "Quant Analyst", "2024-11-28", 20, "$92,000"],
            ["CloudScale Labs", "Cloud/DevOps", "Site Reliability Eng", "2024-12-04", 18, "$88,000"],
            ["CyberGuard Solutions", "Security", "Security Analyst", "2024-12-10", 12, "$80,000"]
        ]
    )

    # 9. CS101_Syllabus.docx
    create_simple_docx(
        os.path.join(sample_root, "CS101_Syllabus.docx"),
        "CS101: Introduction to Computer Systems & Programming Course Syllabus",
        [
            "Instructor: Dr. Alan Turing | Office: Tech Bldg 302",
            "Course Overview: This course covers computer architecture fundamentals, algorithmic problem solving, and Python programming.",
            "Grading Breakdown: Midterm Exam (25%), Final Exam (40%), Assignments (20%), Lab Practicals (15%).",
            "Required Text: Structure and Interpretation of Computer Programs."
        ]
    )

    # 10. Tuition_Challan_Fall2024.pdf
    create_simple_pdf(
        os.path.join(sample_root, "Tuition_Challan_Fall2024.pdf"),
        "COLLEGE BANK FEE CHALLAN - STUDENT COPY",
        [
            "Bank: National Union Bank | College Account No: 882900192841",
            "Challan Number: CHL-FALL2024-551 | Due Date: 30-Aug-2024",
            "Student Name: Alex Morgan | Course: B.Tech Computer Science",
            "Tuition Fees: $4,500.00 | Lab Dues: $350.00",
            "Bank Seal & Teller Signature: Verified & Paid in Cash"
        ]
    )

    return sample_root
